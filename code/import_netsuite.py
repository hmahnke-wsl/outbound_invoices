"""Import NetSuite's invoice and payment history exports: every invoice, what was applied to it, and when.

Two saved searches, exported as .xlsx (or .csv):

- **Invoice History**: one row per invoice, with NetSuite's status ("Open" / "Paid In Full"),
  the amount in its own currency and in USD.
- **Payment History**: one row per payment, credit memo, currency revaluation or journal applied
  to an invoice. Its USD `Amount` is what reconciles: the lines for an invoice net to its USD amount.
  ("Amount (Foreign Currency)" is the invoice's total, not the amount applied, so it isn't read.)

History starts at HISTORY_START, the first day the payment export covers. An earlier invoice is
only kept when it was still open then: NetSuite shows it open, or a payment was applied to it since.

Where NetSuite's status and the applied payments disagree, the invoice goes to Needs attention
with the reason (see `db.reconcile_invoice`).

    uv run python code/import_netsuite.py [invoice_history.xlsx payment_history.xlsx]
"""

import csv
import re
import sys
import warnings
from datetime import date, datetime
from pathlib import Path

from db import DEFAULT_DB, init_db, reconcile_invoice, record_refresh, save_application, save_netsuite_invoice
from import_pdfs import DEFAULT_INBOX
from invoice_parser import normalize_customer

HISTORY_START = "2026-01-01"

INVOICE_COLUMNS = {
    "invoice_number": "Document Number",
    "invoice_date": "Date",
    "due_date": "Due Date/Receive By",
    "netsuite_status": "Status",
    "customer_raw": "Name",
    "subsidiary": "Subsidiary (no hierarchy)",
    "amount_usd": "Amount (Debit)",
    "currency": "Currency",
    "amount": "Amount (Debit) (Foreign Currency)",
    "exchange_rate": "Exchange Rate",
    "memo": "Memo",
}
PAYMENT_COLUMNS = {
    "invoice_number": "Transaction Number",
    "document_number": "Document Number",
    "kind": "Payment Type",
    "applied_on": "Date",
    "amount": "Amount",
    "currency": "Currency: Currency Symbol",
}

PAYMENT_KINDS = {"payment": "payment", "credit memo": "credit_memo", "currency revaluation": "revaluation",
                 "journal": "journal"}
NETSUITE_STATUSES = {"open": "Open", "paid in full": "Paid In Full"}
# NetSuite's currency names, as the invoice export spells them.
CURRENCIES = {"us dollar": "USD", "australian dollar": "AUD", "france euro": "EUR", "euro": "EUR",
              "british pound": "GBP", "south africa rand": "ZAR", "brazilian real": "BRL",
              "canadian dollar": "CAD", "new zealand dollar": "NZD", "japanese yen": "JPY"}


def _cells(path):
    """Every row of the file's first sheet, as tuples of cell values."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as f:
            yield from (tuple(row) for row in csv.reader(f))
        return
    from openpyxl import load_workbook

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # NetSuite's workbooks have no default style; openpyxl says so
        workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        yield from workbook.worksheets[0].iter_rows(values_only=True)
    finally:
        workbook.close()


def read_table(path, columns, key):
    """Rows as dicts named by `columns` ({field: header}), from the first row holding every header.

    Report title rows above the header and total rows below the data (no `key`) are skipped.
    """
    rows = _cells(path)
    for row in rows:
        headers = [str(c).strip() if c is not None else "" for c in row]
        if all(h in headers for h in columns.values()):
            index = {field: headers.index(header) for field, header in columns.items()}
            break
    else:
        raise ValueError(f"{Path(path).name}: no header row with {', '.join(columns.values())}")
    for row in rows:
        record = {field: row[i] if i < len(row) else None for field, i in index.items()}
        record = {f: (v.strip() or None) if isinstance(v, str) else v for f, v in record.items()}
        if record[key]:
            yield record


def _date(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"not a date: {value!r}")


def _number(value):
    if value is None or isinstance(value, (int, float)):
        return None if value is None else float(value)
    text = re.sub(r"[,$\s]", "", value)
    if text.startswith("(") and text.endswith(")"):  # accounting negative
        text = "-" + text[1:-1]
    return float(text) if text else None


def _currency(name):
    if not name:
        return None
    return CURRENCIES.get(name.lower()) or (name.upper() if re.fullmatch(r"[A-Za-z]{3}", name) else name)


def _subsidiary(name):
    match = re.match(r"\s*(\d+)\b", name or "")
    return match.group(1) if match else name


def read_invoices(path):
    """NetSuite's invoice export as invoice dicts (see `db.save_netsuite_invoice`)."""
    for r in read_table(path, INVOICE_COLUMNS, key="invoice_number"):
        status = NETSUITE_STATUSES.get((r["netsuite_status"] or "").lower())
        if status is None:
            raise ValueError(f"{r['invoice_number']}: unknown NetSuite status {r['netsuite_status']!r}")
        yield {
            "invoice_number": r["invoice_number"],
            "customer": normalize_customer(r["customer_raw"]),
            "customer_raw": r["customer_raw"],
            "subsidiary": _subsidiary(r["subsidiary"]),
            "invoice_date": _date(r["invoice_date"]),
            "due_date": _date(r["due_date"]),
            "amount": _number(r["amount"]),
            "currency": _currency(r["currency"]),
            "amount_usd": _number(r["amount_usd"]),
            "exchange_rate": _number(r["exchange_rate"]),
            "memo": r["memo"],
            "netsuite_status": status,
        }


def read_applications(path):
    """NetSuite's payment export as application dicts (see `db.save_application`).

    Its `Amount` is negative where it reduces the invoice; stored the other way round.
    """
    for r in read_table(path, PAYMENT_COLUMNS, key="invoice_number"):
        kind = PAYMENT_KINDS.get((r["kind"] or "").lower())
        if kind is None:
            raise ValueError(f"{r['document_number']}: unknown payment type {r['kind']!r}")
        yield {
            "invoice_number": r["invoice_number"],
            "document_number": r["document_number"],
            "kind": kind,
            "applied_on": _date(r["applied_on"]),
            "amount_usd": -_number(r["amount"]),
            "currency": _currency(r["currency"]),
            "source": Path(path).name,
        }


def import_exports(conn, invoices_path, payments_path, *, start=HISTORY_START, now=None):
    """Import both exports in one transaction. Returns counts of what changed and what needs attention."""
    now = now or datetime.now().astimezone().isoformat(timespec="seconds")
    applications = list(read_applications(payments_path))
    invoices = list(read_invoices(invoices_path))
    applied = {a["invoice_number"] for a in applications}
    applied |= {r[0] for r in conn.execute("SELECT DISTINCT invoice_number FROM applications")}
    known = {r[0] for r in conn.execute("SELECT invoice_number FROM invoices WHERE invoice_number IS NOT NULL")}

    result = {"invoices": 0, "new_invoices": 0, "skipped_before_start": 0, "new_applications": 0,
              "unmatched_applications": 0, "needs_attention": 0}
    with conn:
        touched = {}  # invoice id -> whether NetSuite's status should be applied
        for invoice in invoices:
            number = invoice["invoice_number"]
            if invoice["invoice_date"] and invoice["invoice_date"] < start and invoice["netsuite_status"] != "Open" \
                    and number not in applied and number not in known:
                result["skipped_before_start"] += 1
                continue
            invoice_id, status_changed = save_netsuite_invoice(conn, invoice, now=now)
            touched[invoice_id] = status_changed
            result["invoices"] += 1
            result["new_invoices"] += number not in known
        for application in applications:
            if save_application(conn, application):
                result["new_applications"] += 1
                row = conn.execute("SELECT id FROM invoices WHERE invoice_number = ?",
                                   (application["invoice_number"],)).fetchone()
                if row:
                    touched.setdefault(row["id"], False)
                else:
                    result["unmatched_applications"] += 1
        for invoice_id, status_changed in touched.items():
            reconcile_invoice(conn, invoice_id, apply_status=status_changed, now=now)
        if touched:
            marks = ", ".join("?" * len(touched))
            result["needs_attention"] = conn.execute(
                f"SELECT COUNT(*) FROM invoices WHERE id IN ({marks}) AND ignored = 0 "
                "AND (needs_review = 1 OR reconcile_issue IS NOT NULL)", tuple(touched)).fetchone()[0]
    return result


def _latest(pattern):
    found = sorted(Path(DEFAULT_INBOX).glob(pattern), key=lambda p: p.stat().st_mtime)
    if not found:
        sys.exit(f"no {pattern} in {DEFAULT_INBOX}; pass the two export files")
    return found[-1]


def main():
    if len(sys.argv) == 3:
        invoices_path, payments_path = sys.argv[1:]
    elif len(sys.argv) == 1:
        invoices_path, payments_path = _latest("Invoice History*"), _latest("Payment History*")
    else:
        sys.exit(__doc__.strip().splitlines()[-1].strip())
    print(f"Invoices: {Path(invoices_path).name}\nPayments: {Path(payments_path).name}")
    conn = init_db(DEFAULT_DB)
    result = import_exports(conn, invoices_path, payments_path)
    record_refresh(conn, datetime.now().astimezone().isoformat(timespec="seconds"))
    print(f"{result['invoices']} invoices ({result['new_invoices']} new, {result['skipped_before_start']} paid before "
          f"{HISTORY_START} skipped), {result['new_applications']} new payment lines "
          f"({result['unmatched_applications']} for invoices not in the tracker), "
          f"{result['needs_attention']} need attention")


if __name__ == "__main__":
    main()
