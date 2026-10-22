"""Tests for code/import_netsuite.py: NetSuite's invoice and payment exports into the tracker.

The fixture rows are modelled on real lines from the Jan 2025 - Oct 2026 exports (same layout,
same figures where they illustrate a case) and written to workbooks shaped like NetSuite's.
"""

import csv
import sqlite3
from datetime import datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from api import create_app
from db import init_db, list_needs_attention, list_open, list_paid, save_parsed_invoice, update_invoice
from import_netsuite import import_exports
from invoice_parser import parse_invoice_text
from test_invoice_parser import _usd_invoice

NOW = "2026-10-22T09:00:00-07:00"
LATER = "2026-10-23T09:00:00-07:00"

INVOICE_HEADER = ("Subsidiary (no hierarchy)", "Document Number", "Date", "Due Date/Receive By", "Period", "Status",
                  "Name", "Type", "Account Type", "Amount (Debit)", "Currency", "Amount (Debit) (Foreign Currency)",
                  "Exchange Rate", "Amount (Credit) (Foreign Currency)", "Amount (Credit)", "Memo")
PAYMENT_HEADER = ("Subsidiary: Name", "Entity", "Document Number", "Transaction", "Payment Type", "Date",
                  "Payment Method", "Check Number", "Amount", "Amount (Foreign Currency)", "Currency: Currency Symbol",
                  "Transaction Number")
LLC, MERCH = "210 - WSL, LLC", "315 - WSL Merchandise"


def _day(text):
    return datetime.fromisoformat(text) if text else None


def invoice(number, date, status, name, usd, currency="US Dollar", amount=None, rate=1.0, due=None, memo=None,
            subsidiary=LLC):
    amount = usd if amount is None and currency == "US Dollar" else amount
    return (subsidiary, number, _day(date), _day(due or date), "", status, name, "Invoice", "Accounts Receivable",
            usd, currency, amount, rate, None, None, memo)


def payment(number, document, kind, date, usd, currency="USD", customer="Customer"):
    """`usd` as NetSuite exports it: negative reduces the invoice."""
    return (LLC, customer, document, f"Invoice #{number} - {customer}", kind, _day(date), None, None, usd, None,
            currency, number)


INVOICES = [
    # 2025, paid during 2025: before the history starts.
    invoice("INV4750", "2025-01-01", "Paid In Full", "AB InBev NV/SA", 999500.0),
    # 2025, still open, part paid in 2026.
    invoice("INV4983", "2025-01-01", "Open", "Florence Marine X - 210", 80000.0, due="2025-01-28"),
    # 2025, paid in 2026, in euros: a currency revaluation plus the payment net to its USD amount.
    invoice("INV5880", "2025-12-01", "Paid In Full", "Montres TUDOR SA_EUR", 319948.75, "France Euro", 275000.0, 1.16345),
    invoice("INV6115", "2026-01-01", "Paid In Full", "Team One", 133960.0),  # settled by a credit memo
    invoice("INV6552", "2026-06-01", "Paid In Full", "True Axis", 100000.0),  # $10 short, journalled off
    # One payment covering two invoices; a cent of revaluation booked before it.
    invoice("INV6780", "2026-06-15", "Paid In Full", "OMD Digital_AU", 50000.01, "Australian Dollar", 78295.5, 0.6386),
    invoice("INV6800", "2026-07-01", "Paid In Full", "OMD Digital_AU", 20000.0, "Australian Dollar", 31300.0, 0.639),
    # Merchandise: paid in NetSuite, but the export has no merchandise payments.
    invoice("INV6548", "2026-05-22", "Paid In Full", "Jack's Surfboards", 214120.0, subsidiary=MERCH),
    invoice("INV6693", "2026-06-30", "Paid In Full", "OMD Digital", None),  # no amount
    invoice("INV6999", "2026-09-01", "Open", "Aqua Elite LLC", 12356.0),  # paid, but NetSuite still open
    invoice("INV7025", "2026-10-01", "Paid In Full", "Montres TUDOR SA_EUR", 329834.38, "France Euro", 287500.0, 1.14725,
            due="2026-11-01", memo="2026/2027 Nazare Sponsorship (1 of 2)"),
    invoice("INV7074", "2026-10-19", "Open", "FRANCE TELEVISIONS, SA", 30000.0),
]
PAYMENTS = [
    payment("INV4983", "PYMT7200", "Payment", "2026-09-08", -30000),
    payment("INV5880", "CR1857", "Currency Revaluation", "2026-01-26", 909.58),
    payment("INV5880", "PYMT6462", "Payment", "2026-01-26", -320858.33, "EUR"),
    payment("INV6115", "CM515", "Credit Memo", "2026-01-01", -133960),
    payment("INV6552", "PYMT7092", "Payment", "2026-08-04", -99990),
    payment("INV6552", "JE45846", "Journal", "2026-08-04", -10),
    payment("INV6780", "CR2075", "Currency Revaluation", "2026-07-09", -0.01),
    payment("INV6780", "PYMT7168", "Payment", "2026-08-17", -50000, "AUD"),
    payment("INV6800", "PYMT7168", "Payment", "2026-08-17", -20000, "AUD"),
    payment("INV6999", "PYMT7315", "Payment", "2026-10-21", -12356),
    payment("INV7025", "CR2216", "Currency Revaluation", "2026-10-19", -7636),
    payment("INV7025", "PYMT7311", "Payment", "2026-10-19", -322198.38, "EUR"),
]


def write_invoices(path, rows):
    book = Workbook()
    book.active.append(INVOICE_HEADER)
    for row in rows:
        book.active.append(row)
    book.save(path)
    return path


def write_payments(path, rows):
    """Laid out as NetSuite's report: title lines above the header, a total line below."""
    book = Workbook()
    sheet = book.active
    for line in ("World Surf League", "Parent : 110 - WSL Holdings", "Custom A/R Payment History",
                 "January 1, 2025 - October 22, 2026", None, "Options: Show Zeros"):
        sheet.append((line,))
    sheet.append(PAYMENT_HEADER)
    for row in rows:
        sheet.append(row)
    sheet.append(("Amount Open", None, None, None, None, None, None, None, f"=SUM(I8:I{7 + len(rows)})"))
    book.save(path)
    return path


@pytest.fixture
def exports(tmp_path):
    return write_invoices(tmp_path / "invoices.xlsx", INVOICES), write_payments(tmp_path / "payments.xlsx", PAYMENTS)


@pytest.fixture
def conn(tmp_path):
    conn = init_db(tmp_path / "invoices.db")
    yield conn
    conn.close()


def _invoice(conn, number):
    return dict(conn.execute("SELECT * FROM invoices WHERE invoice_number = ?", (number,)).fetchone())


def _history(conn, number):
    return [(r["status"], r["source"]) for r in conn.execute(
        "SELECT h.status, h.source FROM status_history h JOIN invoices i ON i.id = h.invoice_id "
        "WHERE i.invoice_number = ? ORDER BY h.id", (number,))]


def test_imports_from_the_history_start(conn, exports):
    result = import_exports(conn, *exports, now=NOW)
    assert result == {"invoices": 11, "new_invoices": 11, "skipped_before_start": 1, "new_applications": 12,
                      "unmatched_applications": 0, "needs_attention": 3}
    numbers = {r[0] for r in conn.execute("SELECT invoice_number FROM invoices")}
    assert "INV4750" not in numbers  # paid in 2025
    assert {"INV4983", "INV5880"} <= numbers  # 2025, but open on 1 Jan 2026


def test_reads_netsuite_fields_and_normalizes_the_customer(conn, exports):
    import_exports(conn, *exports, now=NOW)
    tudor = _invoice(conn, "INV7025")
    assert tudor["customer"] == "Montres TUDOR SA" and tudor["customer_raw"] == "Montres TUDOR SA_EUR"
    assert (tudor["amount"], tudor["currency"], tudor["amount_usd"], tudor["exchange_rate"]) == \
        (287500.0, "EUR", 329834.38, 1.14725)
    assert (tudor["invoice_date"], tudor["due_date"], tudor["subsidiary"]) == ("2026-10-01", "2026-11-01", "210")
    assert tudor["memo"] == "2026/2027 Nazare Sponsorship (1 of 2)"
    assert _invoice(conn, "INV6548")["subsidiary"] == "315"
    assert _invoice(conn, "INV4983")["customer"] == "Florence Marine X"


def test_paid_invoices_are_dated_by_when_their_payments_settled_them(conn, exports):
    import_exports(conn, *exports, now=NOW)
    paid = {r["invoice_number"]: r["paid_at"] for r in list_paid(conn)}
    assert paid == {
        "INV5880": "2026-01-26",  # revaluation + euro payment
        "INV6115": "2026-01-01",  # credit memo
        "INV6552": "2026-08-04",  # payment + $10 journal
        "INV6780": "2026-08-17",  # not the earlier revaluation's date
        "INV6800": "2026-08-17",
        "INV7025": "2026-10-19",
    }
    assert _history(conn, "INV7025") == [("open", "netsuite"), ("paid", "netsuite")]


def test_part_paid_and_unpaid_invoices_stay_open(conn, exports):
    import_exports(conn, *exports, now=NOW)
    assert {r["invoice_number"] for r in list_open(conn)} == {"INV4983", "INV7074"}
    assert _invoice(conn, "INV4983")["reconcile_issue"] is None


def test_mismatches_with_netsuite_need_attention(conn, exports):
    import_exports(conn, *exports, now=NOW)
    reasons = {r["invoice_number"]: r["reason"] for r in list_needs_attention(conn)}
    assert reasons == {
        "INV6548": "paid in full in NetSuite, but the payments export has nothing applied to it",
        "INV6693": "no amount found",
        "INV6999": "the payments export covers it in full, but NetSuite still shows it open",
    }


def test_part_of_a_paid_invoice_missing_from_payments_is_flagged(conn, tmp_path):
    exports = (write_invoices(tmp_path / "i.xlsx", [invoice("INV6183", "2026-03-01", "Paid In Full", "One Sport", 12000.0)]),
               write_payments(tmp_path / "p.xlsx", [payment("INV6183", "PYMT6602", "Payment", "2026-03-24", -5967)]))
    import_exports(conn, *exports, now=NOW)
    issue = _invoice(conn, "INV6183")["reconcile_issue"]
    assert issue == "paid in full in NetSuite, but $6,033.00 of it has no payment recorded"


def test_reimport_adds_nothing(conn, exports):
    import_exports(conn, *exports, now=NOW)
    result = import_exports(conn, *exports, now=LATER)
    assert result["new_invoices"] == 0 and result["new_applications"] == 0
    assert conn.execute("SELECT COUNT(*) FROM applications").fetchone()[0] == 12
    assert _history(conn, "INV7025") == [("open", "netsuite"), ("paid", "netsuite")]
    assert _invoice(conn, "INV7025")["netsuite_seen_at"] == LATER


def test_manual_mark_paid_sticks_until_netsuite_changes(conn, exports, tmp_path):
    import_exports(conn, *exports, now=NOW)
    ftv = _invoice(conn, "INV7074")
    update_invoice(conn, ftv["id"], {"status": "paid", "paid_at": "2026-10-21"}, NOW)
    import_exports(conn, *exports, now=LATER)  # NetSuite still says open
    assert (_invoice(conn, "INV7074")["status"], _invoice(conn, "INV7074")["paid_at"]) == ("paid", "2026-10-21")

    # NetSuite now records the rest of INV4983's payment.
    rows = [invoice("INV4983", "2025-01-01", "Paid In Full", "Florence Marine X - 210", 80000.0, due="2025-01-28")]
    later = (write_invoices(tmp_path / "i2.xlsx", rows),
             write_payments(tmp_path / "p2.xlsx", [payment("INV4983", "PYMT7400", "Payment", "2026-10-22", -50000)]))
    result = import_exports(conn, *later, now=LATER)
    assert result["new_applications"] == 1
    florence = _invoice(conn, "INV4983")
    assert (florence["status"], florence["paid_at"], florence["reconcile_issue"]) == ("paid", "2026-10-22", None)
    assert _history(conn, "INV4983") == [("open", "netsuite"), ("paid", "netsuite")]


def test_netsuite_fills_in_a_pdf_invoice_and_keeps_its_terms_and_po(conn, tmp_path):
    parsed = parse_invoice_text(_usd_invoice(), "acme.pdf") | {"source_file": "acme.pdf"}
    save_parsed_invoice(conn, parsed, pdf_file="acme.pdf", now=NOW)
    rows = [invoice("INV7100", "2026-09-15", "Open", "Acme Beverage Co_210", 42000.0, due="2026-10-15",
                    memo="2026 Event Sponsorship")]
    result = import_exports(conn, write_invoices(tmp_path / "i.xlsx", rows), write_payments(tmp_path / "p.xlsx", []),
                            now=NOW)
    assert result["new_invoices"] == 0
    acme = _invoice(conn, "INV7100")
    assert (acme["customer"], acme["terms"], acme["po_number"], acme["pdf_file"]) == \
        ("Acme Beverage Co", "Net 30", "4500123", "acme.pdf")
    assert (acme["netsuite_status"], acme["amount_usd"], acme["status"]) == ("Open", 42000.0, "open")
    assert _history(conn, "INV7100") == [("open", "import")]


def test_netsuite_does_not_overwrite_a_manual_fix(conn, tmp_path):
    flagged = _usd_invoice(amount_line="Amount Due")
    save_parsed_invoice(conn, parse_invoice_text(flagged, "acme.pdf") | {"source_file": "acme.pdf"}, now=NOW)
    update_invoice(conn, 1, {"amount": 41000}, NOW)
    rows = [invoice("INV7100", "2026-09-15", "Open", "Acme Beverage Co", 42000.0)]
    import_exports(conn, write_invoices(tmp_path / "i.xlsx", rows), write_payments(tmp_path / "p.xlsx", []), now=NOW)
    assert _invoice(conn, "INV7100")["amount"] == 41000


def test_netsuite_clears_a_flag_the_pdf_left(conn, tmp_path):
    flagged = _usd_invoice(amount_line="Amount Due")
    save_parsed_invoice(conn, parse_invoice_text(flagged, "acme.pdf") | {"source_file": "acme.pdf"}, now=NOW)
    assert _invoice(conn, "INV7100")["needs_review"] == 1
    rows = [invoice("INV7100", "2026-09-15", "Open", "Acme Beverage Co", 42000.0)]
    import_exports(conn, write_invoices(tmp_path / "i.xlsx", rows), write_payments(tmp_path / "p.xlsx", []), now=NOW)
    assert list_needs_attention(conn) == []
    assert [r["amount"] for r in list_open(conn)] == [42000.0]


def test_flagged_by_netsuite_cannot_be_marked_paid_but_can_be_ignored(conn, exports):
    import_exports(conn, *exports, now=NOW)
    merch = _invoice(conn, "INV6548")
    with pytest.raises(ValueError):
        update_invoice(conn, merch["id"], {"status": "paid", "paid_at": "2026-06-01"}, NOW)
    update_invoice(conn, merch["id"], {"ignored": True}, NOW)
    assert "INV6548" not in {r["invoice_number"] for r in list_needs_attention(conn)}


def test_payment_for_an_invoice_not_yet_imported_is_kept(conn, tmp_path):
    result = import_exports(conn, write_invoices(tmp_path / "i.xlsx", []),
                            write_payments(tmp_path / "p.xlsx", [payment("INV7200", "PYMT1", "Payment", "2026-10-20", -500)]),
                            now=NOW)
    assert result["unmatched_applications"] == 1
    import_exports(conn, write_invoices(tmp_path / "i2.xlsx", [invoice("INV7200", "2026-10-01", "Paid In Full", "X", 500.0)]),
                   write_payments(tmp_path / "p2.xlsx", []), now=LATER)
    assert (_invoice(conn, "INV7200")["status"], _invoice(conn, "INV7200")["paid_at"]) == ("paid", "2026-10-20")


def test_reads_csv_exports(conn, tmp_path):
    invoices = tmp_path / "invoices.csv"
    with invoices.open("w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(INVOICE_HEADER)
        writer.writerow([LLC, "INV7074", "10/19/2026", "10/19/2026", "Oct 2026", "Paid In Full", "FRANCE TELEVISIONS, SA",
                         "Invoice", "Accounts Receivable", "30,000.00", "US Dollar", "30,000.00", "1", "", "", "Rights"])
    payments = tmp_path / "payments.csv"
    with payments.open("w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["Custom A/R Payment History"])
        writer.writerow(PAYMENT_HEADER)
        writer.writerow([LLC, "FTV", "PYMT1", "", "Payment", "2026-10-21", "", "", "(30,000.00)", "", "USD", "INV7074"])
    import_exports(conn, invoices, payments, now=NOW)
    ftv = _invoice(conn, "INV7074")
    assert (ftv["status"], ftv["paid_at"], ftv["amount"], ftv["reconcile_issue"]) == ("paid", "2026-10-21", 30000.0, None)


def test_rejects_a_file_without_the_expected_columns(conn, tmp_path):
    book = Workbook()
    book.active.append(("Document Number", "Date"))
    book.save(tmp_path / "wrong.xlsx")
    with pytest.raises(ValueError, match="no header row"):
        import_exports(conn, tmp_path / "wrong.xlsx", write_payments(tmp_path / "p.xlsx", []), now=NOW)


def test_upgrades_a_database_made_before_netsuite(tmp_path):
    path = tmp_path / "old.db"
    old = sqlite3.connect(path)
    old.executescript("""
        CREATE TABLE invoices (id INTEGER PRIMARY KEY, customer TEXT NOT NULL, invoice_number TEXT, invoice_date TEXT,
            terms TEXT, due_date TEXT, amount REAL, currency TEXT NOT NULL DEFAULT 'USD', memo TEXT, po_number TEXT,
            status TEXT NOT NULL DEFAULT 'open', paid_at TEXT, needs_review INTEGER NOT NULL DEFAULT 0,
            review_reason TEXT, source_file TEXT, pdf_file TEXT, source_link TEXT, sent_at TEXT,
            manually_edited INTEGER NOT NULL DEFAULT 0, ignored INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE status_history (id INTEGER PRIMARY KEY, invoice_id INTEGER NOT NULL REFERENCES invoices(id),
            status TEXT NOT NULL, changed_at TEXT NOT NULL, source TEXT NOT NULL CHECK (source IN ('import', 'manual')));
        INSERT INTO invoices (customer, invoice_number) VALUES ('Acme', 'INV1');
        INSERT INTO status_history (invoice_id, status, changed_at, source) VALUES (1, 'open', '2026-10-01T00:00:00+00:00', 'import');
    """)
    old.close()
    conn = init_db(path)
    conn.execute("INSERT INTO status_history (invoice_id, status, changed_at, source) VALUES (1, 'paid', ?, 'netsuite')", (NOW,))
    assert _history(conn, "INV1") == [("open", "import"), ("paid", "netsuite")]
    assert _invoice(conn, "INV1")["reconcile_issue"] is None
    conn.close()


def test_api_history_shows_netsuite_as_the_source(tmp_path, exports):
    path = tmp_path / "api.db"
    conn = init_db(path)
    import_exports(conn, *exports, now=NOW)
    tudor = _invoice(conn, "INV7025")["id"]
    conn.close()
    client = TestClient(create_app(path, today="2026-10-22", now=NOW, attachments_dir=tmp_path))
    assert [h["source"] for h in client.get(f"/api/invoices/{tudor}/history").json()] == ["netsuite", "netsuite"]
    assert len(client.get("/api/needs-attention").json()) == 3
